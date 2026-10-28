"""Build the recorded university run from the archived research materials.

The source folder is intentionally outside this repository. Pass its path with
``--source``. The output is a read-only ``sjt-run/1`` recording: it contains
the review-agent issues and the 0/1/2 scores that are present in the archive,
and never invents expert scores or missing model parameters.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

DIMENSIONS = ["construct", "narrative", "fairness", "realism"]
DIMENSION_LABELS = {
    "construct": "构念一致性",
    "narrative": "叙事衔接",
    "fairness": "公平无偏",
    "realism": "现实可行",
}
RUN_ID = "recorded-univ-echoes-2.0"
OUTLINE_VERSION = "ov4"
STORY_VERSION = "v4"
ITEM_VERSION = "v4"
SCORING_RULE = "zero_or_two_ones_same_dimension"


def utc(value: str) -> str:
    return value if value.endswith("Z") else value + "Z"


def clean_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value)).strip()


def score(value: object) -> int | None:
    if value is None or value == "":
        return None
    try:
        number = float(str(value).replace("**", "").strip())
    except ValueError:
        return None
    if number not in (0, 1, 2):
        return None
    return int(number)


# ---------------------------------------------------------------- raw sources


def parse_source_items(outline_docx: Path) -> list[dict]:
    doc = Document(outline_docx)
    lines = [clean_text(p.text) for p in doc.paragraphs if clean_text(p.text)]
    records: dict[int, dict] = {}
    for line in lines:
        number = re.match(r"^(\d+)\s*\{", line)
        if not number:
            continue
        text = re.search(r'"original_item"\s*:\s*"([^"]*)"', line)
        facet = re.search(r'"dimension"\s*:\s*"([^"]*)"', line)
        if not text or not facet:
            continue
        records[int(number.group(1))] = {
            "order": int(number.group(1)),
            "text": text.group(1),
            "facet": facet.group(1),
        }
    if sorted(records) != list(range(1, 61)):
        raise ValueError("outline docx must contain 60 source-item records")
    return [records[order] for order in range(1, 61)]


def parse_outline_events(outline_docx: Path) -> list[str]:
    doc = Document(outline_docx)
    lines = [clean_text(p.text) for p in doc.paragraphs if clean_text(p.text)]
    start = next(i for i, line in enumerate(lines) if line.startswith("Day 1"))
    end = next(i for i, line in enumerate(lines) if line.startswith("按序排列"))
    events: list[str] = []
    pattern = re.compile(
        r"^(?:\d+\.\s*)?(.+?)\s*\((Day\s+\d+\s+[^)]+)\)\s*-\s*(.+)$"
    )
    for line in lines[start:end]:
        match = pattern.match(line)
        if not match:
            continue
        _dimension, when, description = match.groups()
        events.append(f"{when.strip()} · {description.strip()}")
    if len(events) != 60:
        raise ValueError(f"expected 60 outline events, got {len(events)}")
    return events


SCENARIO_PATTERN = re.compile(
    r"\[情景(\d+)\]\s*正文：\s*(.*?)\s*问句：\s*(.*?)\s*选项：\s*(.*?)(?=\n\s*\[情景|\Z)",
    re.S,
)


def parse_story_file(path: Path, with_levels: bool = False) -> dict[int, dict]:
    text = path.read_text(encoding="utf-8-sig")
    scenarios: dict[int, dict] = {}
    for match in SCENARIO_PATTERN.finditer(text):
        number = int(match.group(1))
        options_block = match.group(4).strip()
        options: dict[str, str] = {}
        levels: dict[str, int | None] = {}
        for option_match in re.finditer(
            r"^\s*([A-D])[.、．]\s*(.+?)\s*$", options_block, re.M
        ):
            letter = option_match.group(1)
            value = clean_text(option_match.group(2))
            level: int | None = None
            if with_levels:
                level_match = re.search(r"[（(]\s*([1-4])\s*分\s*[）)]\s*$", value)
                if level_match:
                    level = int(level_match.group(1))
                    value = value[: level_match.start()].strip()
            options[letter] = value
            levels[letter] = level
        if set(options) != set("ABCD"):
            raise ValueError(f"{path}: scenario {number} needs A-D options")
        if with_levels and sorted(
            level for level in levels.values() if level is not None
        ) != [1, 2, 3, 4]:
            raise ValueError(f"{path}: scenario {number} needs option levels 1-4")
        scenarios[number] = {
            "text": clean_text(match.group(2)),
            "question": clean_text(match.group(3)),
            "options": options,
            "levels": levels,
        }
    if sorted(scenarios) != list(range(1, 61)):
        raise ValueError(f"{path}: expected 60 scenarios, got {len(scenarios)}")
    return scenarios


def story_only(scenarios: dict[int, dict]) -> dict[int, dict]:
    return {
        number: {
            "text": value["text"],
            "question": value["question"],
            "options": value["options"],
        }
        for number, value in scenarios.items()
    }


def scenario_block(scenario: dict) -> str:
    options = "\n".join(
        f"{letter}. {scenario['options'][letter]}" for letter in "ABCD"
    )
    return f"正文：{scenario['text']}\n问句：{scenario['question']}\n选项：\n{options}"


def field_text(scenario: dict, field: str) -> str:
    if field == "story":
        return scenario["text"]
    if field == "question":
        return scenario["question"]
    if field.startswith("option:"):
        return scenario["options"][field.split(":")[1]]
    raise ValueError(f"unknown field {field}")


# ---------------------------------------------------------------- review agent


def parse_review_issue_file(path: Path) -> dict[int, list[dict]]:
    """Return the before/after scenario text pairs in the second review file."""
    text = path.read_text(encoding="utf-8-sig")
    markers = list(re.finditer(r"\[情景(\d+)\]", text))
    blocks: dict[int, list[dict]] = {}
    for index, marker in enumerate(markers):
        number = int(marker.group(1))
        end = markers[index + 1].start() if index + 1 < len(markers) else len(text)
        chunk = text[marker.end() : end]
        # The archive separates paired scenarios with a dashed rule; feedback
        # paragraphs follow the first block but never start with "A. " etc.
        parts = re.split(r"\n\s*[-—]{6,}\s*\n", chunk)
        for part in parts:
            body = part.strip()
            if "正文" not in body or "问句" not in body or "选项" not in body:
                continue
            try:
                story_match = re.search(r"正文：\s*(.*?)\s*问句：", body, re.S)
                question_match = re.search(r"问句：\s*(.*?)\s*选项：", body, re.S)
                options_part = body.split("选项：", 1)[1]
            except IndexError:
                continue
            options: dict[str, str] = {}
            for option_match in re.finditer(
                r"^\s*([A-D])[.、．]\s*(.+?)\s*$", options_part, re.M
            ):
                options.setdefault(option_match.group(1), clean_text(option_match.group(2)))
            if set(options) != set("ABCD"):
                continue
            blocks.setdefault(number, []).append(
                {
                    "text": clean_text(story_match.group(1)),
                    "question": clean_text(question_match.group(1)),
                    "options": options,
                }
            )
    return {number: values[:2] for number, values in blocks.items() if len(values) >= 2}


GEMINI_ISSUES = [
    {
        "scenario": 5,
        "target_stage": "story_generation",
        "category": "逻辑冲突",
        "problem": "正文写“手边的咖啡还冒着热气”，选项 A 却是“下楼买咖啡”，同一时刻自相矛盾。",
        "field": "story",
    },
    {
        "scenario": 9,
        "target_stage": "item_transformation",
        "category": "常识错误",
        "problem": "样本需要冷藏过夜，选项 C 却让受测者把样本装进书包带走。",
        "field": "option:C",
    },
    {
        "scenario": 17,
        "target_stage": "story_generation",
        "category": "空间冲突",
        "problem": "正文称对方为“室友”，选项里却要求“送到对方房间”，居住关系自相矛盾。",
        "field": "story",
    },
    {
        "scenario": 27,
        "target_stage": "item_transformation",
        "category": "可行性问题",
        "problem": "距开考只剩五分钟，选项 D 却“按平稳步伐前行”，冷静被写成了必然迟到。",
        "field": "option:D",
    },
    {
        "scenario": 4,
        "target_stage": "item_transformation",
        "category": "语体问题",
        "problem": "问句使用“视觉刺激”这一临床化表述，与叙事语体脱节。",
        "field": "question",
    },
    {
        "scenario": 7,
        "target_stage": "item_transformation",
        "category": "表达过度",
        "problem": "选项 A“猛拍桌子又大笑”像戏剧角色，不像真实情绪失控。",
        "field": "option:A",
    },
]

SECOND_REVIEW_ISSUES = [
    {
        "scenario": 2,
        "target_stage": "item_transformation",
        "category": "构念偏移",
        "problem": "目标测同情（情感），问句“你会怎么做”却在测帮助行为；替人刷校园卡也不符合实名规则。",
    },
    {
        "scenario": 9,
        "target_stage": "item_transformation",
        "category": "构念偏移",
        "problem": "目标测“负责”，剧情和选项却在测计划性；样本冷藏与随手带走的设定仍然冲突。",
    },
    {
        "scenario": 11,
        "target_stage": "item_transformation",
        "category": "构念偏移",
        "problem": "目标测同情，选项却在测是否离座接水等帮助行为，还混入课堂纪律与卫生风险。",
    },
    {
        "scenario": 15,
        "target_stage": "item_transformation",
        "category": "构念偏移",
        "problem": "目标测信任，选项却混入议价、验货等谨慎行为；“事后再联系解决”也缺少现实保障。",
    },
    {
        "scenario": 22,
        "target_stage": "item_transformation",
        "category": "构念污染",
        "problem": "目标测信任，选项却混入“语气生硬”“冷淡”等谦恭维度，以及团队效率线索。",
    },
    {
        "scenario": 31,
        "target_stage": "item_transformation",
        "category": "构念污染",
        "problem": "目标测信任，选项却用“严厉质疑”“板着脸”测到了对抗与礼貌，超出信任判断。",
    },
]


def gemini_review_count(path: Path) -> int:
    text = path.read_text(encoding="utf-8-sig")
    layers: dict[str, dict[int, int]] = {"story": {}, "question": {}, "option": {}}
    current: str | None = None
    for line in text.splitlines():
        if "### 1. 剧情审查结果" in line:
            current = "story"
            continue
        if "### 2. 问题审查结果" in line:
            current = "question"
            continue
        if "### 3. 选项审查结果" in line:
            current = "option"
            continue
        if "### 4. 总结与风险评估" in line:
            current = None
            continue
        if current is None or not line.strip().startswith("|"):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if len(cells) < 2:
            continue
        id_text = cells[0].replace("*", "").strip()
        value = score(cells[1].replace("*", ""))
        if value is None:
            continue
        range_match = re.fullmatch(r"(\d+)\s*[-–]\s*(\d+)", id_text)
        if range_match:
            start, end = (int(part) for part in range_match.groups())
            for number in range(start, end + 1):
                layers[current][number] = value
        elif id_text.isdigit():
            layers[current][int(id_text)] = value
    low = {
        number
        for values in layers.values()
        for number, value in values.items()
        if value < 2
    }
    return len(low)


# ---------------------------------------------------------------- scoring rounds


def read_story_scores_merged(
    path: Path, models: list[str], mode: str
) -> dict[str, dict[int, dict[str, int | None]]]:
    """Read the story sheet, keeping only the primary 0/1/2 item scores."""
    wb = load_workbook(path, read_only=True, data_only=True)
    worksheet = wb[wb.sheetnames[0]]
    result = {model: {number: {} for number in range(1, 61)} for model in models}
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        match = re.search(r"(\d+)", str(row[0]))
        if not match:
            continue
        number = int(match.group(1))
        for model_index, model in enumerate(models):
            result[model][number] = {
                dimension: score(
                    row[
                        (1 + 5 * model_index + index)
                        if mode == "permodel"
                        else (1 + 4 * index + model_index)
                    ]
                )
                for index, dimension in enumerate(DIMENSIONS)
            }
    if any(
        len([value for value in result[model][number].values() if value is not None]) == 0
        for model in models
        for number in range(1, 61)
    ):
        raise ValueError(f"{path}: incomplete story score sheet")
    return result


def trigger_reasons(
    scores: dict[str, dict[int, dict[str, int | None]]], number: int
) -> list[str]:
    reasons: list[str] = []
    for dimension in DIMENSIONS:
        values = [
            scores[model].get(number, {}).get(dimension) for model in scores
        ]
        if any(value == 0 for value in values if value is not None):
            zero_models = [
                model
                for model in scores
                if scores[model].get(number, {}).get(dimension) == 0
            ]
            reasons.append(
                f"{DIMENSION_LABELS[dimension]} 0 分（{'、'.join(zero_models)}）"
            )
            continue
        ones = [
            model
            for model in scores
            if scores[model].get(number, {}).get(dimension) == 1
        ]
        if len(ones) >= 2:
            reasons.append(
                f"{DIMENSION_LABELS[dimension]} 同维度 2 个 1 分（{'、'.join(ones)}）"
            )
    return reasons


def build_scoring_record(
    *,
    review_id: str,
    slot: str,
    model_label: str,
    attempt: int,
    version: str,
    own_scores: dict[int, dict[str, int | None]],
    all_scores: dict[str, dict[int, dict[str, int | None]]],
    source_files: list[str],
) -> dict:
    item_scores = {
        f"u-item-{number:02d}": {
            dimension: value
            for dimension, value in own_scores[number].items()
            if value is not None
        }
        for number in range(1, 61)
    }
    triggers = {
        number: trigger_reasons(all_scores, number) for number in range(1, 61)
    }
    triggers = {number: reasons for number, reasons in triggers.items() if reasons}
    low = [
        number
        for number in range(1, 61)
        if any(value < 2 for value in own_scores[number].values() if value is not None)
    ]
    decision = "revise" if triggers else "pass"
    rationale = (
        f"第 {attempt} 轮打分：本轮共 {len(triggers)} 个触发项；"
        f"本模型有 {len(low)} 个条目低于 2 分。"
    )
    issues = []
    if triggers:
        display = "；".join(
            f"情景{number:02d} " + "、".join(reasons)
            for number, reasons in list(triggers.items())[:6]
        )
        issues.append(display)
        if len(triggers) > 6:
            issues.append(f"其余 {len(triggers) - 6} 个触发项见工作簿。")
    else:
        issues.append("本轮无触发项。")
    return {
        "review_id": review_id,
        "reviewer_slot": slot,
        "model_label": model_label,
        "status": "completed",
        "attempt": attempt,
        "revision": version,
        "rubric_version": "ai-eval-0-1-2",
        "reviewed_item_version": version,
        "item_scores": item_scores,
        "recommendation": decision,
        "rationale": rationale,
        "issues": issues,
        "source_files": source_files,
    }


# ------------------------------------------------------------- artifact build


def read_prompt(path: Path, limit: int = 1400) -> str:
    if path.suffix.lower() == ".docx":
        text = "\n".join(
            paragraph.text for paragraph in Document(path).paragraphs if paragraph.text.strip()
        )
    else:
        text = path.read_text(encoding="utf-8-sig")
    text = text.strip()
    if len(text) <= limit:
        return text
    return text[:limit].rstrip() + "\n……（以下省略，完整提示词见研究存档）"


def changed_ids(before: dict[int, dict], after: dict[int, dict]) -> list[int]:
    return [number for number in range(1, 61) if before[number] != after[number]]


def revision_text(scenarios: dict[int, dict], ids: list[int]) -> str:
    return "\n\n".join(
        f"【情景{number:02d}】\n{scenario_block(scenarios[number])}" for number in ids
    )


def build(args: argparse.Namespace) -> dict:
    source = args.source
    source_items = parse_source_items(source / "03大纲设计/大学大纲4.0.docx")
    outline_descriptions = parse_outline_events(source / "03大纲设计/大学大纲4.0.docx")

    story_paths = {
        "v1": source / "04故事生成/大学故事2.0=大纲4.0-gemini.txt",
        "v2": source / "05剧情评价/大学故事2.0=大纲4.0-gemini-改.txt",
        "v3": source / "06AI评价打分/!大学修改后评价/大学故事2.0=大纲4.0-gemini-改.txt",
    }
    stories = {version: story_only(parse_story_file(path)) for version, path in story_paths.items()}
    final_story = parse_story_file(
        source / "07专家评审/大学生剧情修改后最终用版.txt", with_levels=True
    )
    final_scenarios = story_only(final_story)
    stories["v4"] = final_scenarios

    gemini_path = source / "05剧情评价/大学生故事2.0 -gemini.txt"
    issue_path = source / "06AI评价打分/!大学修改后评价/问题情景以及意见大学.txt"
    second_pairs = parse_review_issue_file(issue_path)
    for number in (2, 9, 11, 15, 22, 31):
        if number not in second_pairs:
            raise ValueError(f"second review file is missing scenario {number}")

    review_reports = [
        {
            "report_id": "rr-1",
            "attempt": 1,
            "model_label": "Gemini（具体版本未记录）",
            "status": "completed",
            "reviewed_version": "v1",
            "checked_count": 60,
            "summary": "逐条审查 60 个情景的剧情、问句与选项；23 个情景至少一层低于 2 分。",
            "issues": [
                {
                    "issue_id": f"ri-1-{spec['scenario']:02d}",
                    "scenario_id": f"u-item-{spec['scenario']:02d}",
                    "target_stage": spec["target_stage"],
                    "category": spec["category"],
                    "problem": spec["problem"],
                    "before": field_text(stories["v1"][spec["scenario"]], spec["field"]),
                    "after": field_text(stories["v2"][spec["scenario"]], spec["field"]),
                }
                for spec in GEMINI_ISSUES
            ],
            "source_files": [gemini_path.name],
        },
        {
            "report_id": "rr-2",
            "attempt": 2,
            "model_label": "审查智能体（模型未记录）",
            "status": "completed",
            "reviewed_version": "v2",
            "checked_count": 60,
            "summary": "复审 v2，仍有 6 个情景存在构念或现实可行性问题。",
            "issues": [
                {
                    "issue_id": f"ri-2-{spec['scenario']:02d}",
                    "scenario_id": f"u-item-{spec['scenario']:02d}",
                    "target_stage": spec["target_stage"],
                    "category": spec["category"],
                    "problem": spec["problem"],
                    "before": scenario_block(second_pairs[spec["scenario"]][0]),
                    "after": scenario_block(second_pairs[spec["scenario"]][1]),
                }
                for spec in SECOND_REVIEW_ISSUES
            ],
            "source_files": [issue_path.name],
        },
    ]
    review_reports[0]["summary"] = (
        f"逐条审查 60 个情景的剧情、问句与选项；"
        f"{gemini_review_count(gemini_path)} 个情景至少一层低于 2 分。"
    )

    scoring_sources = [
        {
            "attempt": 1,
            "version": "v3",
            "path": source / "06AI评价打分/大学/合并.xlsx",
            "mode": "interleaved",
            "source_files": ["大学/合并.xlsx"],
            "models": [
                ("G", "review_a", "GPT（具体版本未记录）"),
                ("D", "review_b", "DeepSeek（具体版本未记录）"),
                ("Q", "review_c", "Qwen（具体版本未记录）"),
            ],
        },
        {
            "attempt": 2,
            "version": "v4",
            "path": source / "06AI评价打分/大学问题情节修改后打分/合并.xlsx",
            "mode": "permodel",
            "source_files": ["大学问题情节修改后打分/合并.xlsx"],
            "models": [
                ("G", "review_a", "GPT（具体版本未记录）"),
                ("D", "review_b", "DeepSeek（具体版本未记录）"),
                ("Q", "review_c", "Qwen（具体版本未记录）"),
            ],
        },
    ]

    reviews: list[dict] = []
    round_triggers: dict[int, dict[int, list[str]]] = {}
    for definition in scoring_sources:
        codes = [code for code, _slot, _label in definition["models"]]
        scores = read_story_scores_merged(definition["path"], codes, definition["mode"])
        round_triggers[definition["attempt"]] = {
            number: reasons
            for number in range(1, 61)
            if (reasons := trigger_reasons(scores, number))
        }
        for code, slot, label in definition["models"]:
            reviews.append(
                build_scoring_record(
                    review_id=f"sr-{definition['attempt']}-{slot[-1]}",
                    slot=slot,
                    model_label=label,
                    attempt=definition["attempt"],
                    version=definition["version"],
                    own_scores=scores[code],
                    all_scores=scores,
                    source_files=definition["source_files"],
                )
            )

    if len(round_triggers[1]) != 8 or len(round_triggers[2]) != 0:
        raise ValueError(
            "unexpected scoring triggers: "
            f"round1={len(round_triggers[1])}, round2={len(round_triggers[2])}"
        )

    scoring_examples = [
        {
            "round": 1,
            "scenario_id": "u-item-14",
            "outcome": "revised",
            "reason": "DeepSeek 给“现实可行”0 分；任一 0 分即触发修改。",
            "before": scenario_block(stories["v3"][14]),
            "after": scenario_block(final_scenarios[14]),
        },
        {
            "round": 1,
            "scenario_id": "u-item-33",
            "outcome": "revised",
            "reason": "DeepSeek 给“构念一致性”0 分；任一 0 分即触发修改。",
            "before": scenario_block(stories["v3"][33]),
            "after": scenario_block(final_scenarios[33]),
        },
        {
            "round": 2,
            "scenario_id": "u-item-56",
            "outcome": "passed",
            "reason": "GPT 在三个维度给 1 分，但同一维度没有第二个智能体给 1 分，因此不修改。",
            "before": scenario_block(final_scenarios[56]),
            "after": scenario_block(final_scenarios[56]),
        },
        {
            "round": 2,
            "scenario_id": "u-item-11",
            "outcome": "passed",
            "reason": "Qwen 给“公平无偏”1 分，但只有 1 个智能体给 1 分，因此不修改。",
            "before": scenario_block(final_scenarios[11]),
            "after": scenario_block(final_scenarios[11]),
        },
    ]

    source_records = []
    for number, record in enumerate(source_items, start=1):
        source_records.append(
            {
                "source_item_id": f"src-u{number:02d}",
                "text": record["text"],
                "domain": "大五人格（BFI-2 构念）",
                "facet": record["facet"],
                "reverse_flag": None,
                "note": "原始问卷条目；反向计分标记未在研究材料中记录。",
            }
        )

    outline_events = []
    for number, description in enumerate(outline_descriptions, start=1):
        outline_events.append(
            {
                "event_id": f"oe-{number:02d}",
                "order": number,
                "description": description,
                "source_item_ids": [f"src-u{number:02d}"],
                "item_ids": [f"u-item-{number:02d}"],
            }
        )

    items = []
    for number in range(1, 61):
        scenario = final_story[number]
        items.append(
            {
                "item_id": f"u-item-{number:02d}",
                "version": ITEM_VERSION,
                "source_item_id": f"src-u{number:02d}",
                "item_family_id": f"family-u-{number:02d}",
                "domain": "大五人格（BFI-2 构念）",
                "facet": source_items[number - 1]["facet"],
                "context_id": "university",
                "event_id": f"oe-{number:02d}",
                "stem": scenario["question"],
                "options": [
                    {
                        "option_id": letter,
                        "text": scenario["options"][letter],
                        "candidate_level": scenario["levels"][letter],
                        "reverse": None,
                        "review_flags": [],
                    }
                    for letter in "ABCD"
                ],
                "candidate_scoring_key": "D>C>B>A",
                "expert_confirmed_scoring_key": None,
                "review_markers": [
                    {"reviewer_slot": "review_a", "status": "completed"},
                    {"reviewer_slot": "review_b", "status": "completed"},
                    {"reviewer_slot": "review_c", "status": "completed"},
                ],
                "note": "最终版选项分值 A=1、B=2、C=3、D=4；专家确认计分键待专家评分。",
            }
        )

    revision_specs = [
        (
            "rev_1",
            "v1",
            "v2",
            "审查智能体指出 23 个情景存在逻辑、构念或现实性问题；退回条目与故事阶段修订。",
            ["rr-1"],
        ),
        (
            "rev_2",
            "v2",
            "v3",
            "复审仍有 6 个情景存在构念偏移或可行性问题；按条目意见重写问句与选项。",
            ["rr-2"],
        ),
        (
            "rev_3",
            "v3",
            "v4",
            "三个打分智能体对 v3 给出 8 个触发项（含 0 分）；定向修订后形成最终版 v4。",
            ["sr-1-a", "sr-1-b", "sr-1-c"],
        ),
    ]
    revisions = []
    for revision_id, from_version, to_version, reason, source_ids in revision_specs:
        before = stories[from_version]
        after = stories[to_version]
        ids = changed_ids(before, after)
        revisions.append(
            {
                "revision_id": revision_id,
                "from_version": from_version,
                "to_version": to_version,
                "reason": reason,
                "source_review_ids": source_ids,
                "demonstration_only": False,
                "changed_item_ids": [f"u-item-{number:02d}" for number in ids],
                "before_text": revision_text(before, ids),
                "after_text": revision_text(after, ids),
            }
        )

    events = [
        (1, "2026-01-16T02:00:00Z", "stage_status", "intake", "running", "读取 60 个原始条目。"),
        (2, "2026-01-16T02:05:00Z", "stage_status", "intake", "completed", "60 个原始条目登记完成。"),
        (3, "2026-01-16T02:35:00Z", "stage_status", "item_transformation", "completed", "条目转化为行为化情境目标。"),
        (4, "2026-01-16T02:50:00Z", "stage_status", "outline_generation", "completed", "生成 60 个事件大纲。"),
        (5, "2026-01-16T02:50:00Z", "human_gate", "outline_confirmation", "waiting_human", "大纲确认检查点。"),
        (6, "2026-01-16T03:00:00Z", "stage_status", "outline_confirmation", "completed", "研究者确认大纲顺序。"),
        (7, "2026-01-16T03:20:00Z", "stage_status", "story_generation", "completed", "生成 v1 故事与 60 道四选项题。"),
        (8, "2026-01-16T04:07:00Z", "stage_status", "review_agent", "completed", "审查智能体完成 60 题审查；23 个情景有问题。"),
        (9, "2026-01-16T04:13:00Z", "stage_status", "auto_revision", "completed", "退回条目/故事阶段修订，生成 v2。"),
        (10, "2026-01-16T04:20:00Z", "stage_status", "story_generation", "completed", "v2 生成完成。"),
        (11, "2026-01-23T06:45:00Z", "stage_status", "review_agent", "completed", "复审发现 6 个情景仍有构念或可行性问题。"),
        (12, "2026-01-23T06:48:00Z", "stage_status", "auto_revision", "completed", "按复审意见重写条目，生成 v3。"),
        (13, "2026-01-23T07:00:00Z", "stage_status", "story_generation", "completed", "v3 生成完成。"),
        (14, "2026-01-29T05:45:00Z", "stage_status", "review_a", "completed", "GPT 完成 v3 打分。"),
        (15, "2026-01-29T05:50:00Z", "stage_status", "review_b", "completed", "DeepSeek 完成 v3 打分。"),
        (16, "2026-01-29T06:10:00Z", "stage_status", "review_c", "completed", "Qwen 完成 v3 打分。"),
        (17, "2026-01-29T06:15:00Z", "stage_status", "review_aggregation", "completed", "第 1 轮规则判定：8 个触发项，退回定向修订。"),
        (18, "2026-02-02T01:22:00Z", "stage_status", "auto_revision", "completed", "按打分触发项修订，生成 v4 最终版。"),
        (19, "2026-02-02T01:30:00Z", "stage_status", "story_generation", "completed", "v4 最终版生成完成。"),
        (20, "2026-02-02T02:10:00Z", "stage_status", "review_a", "completed", "GPT 完成 v4 打分。"),
        (21, "2026-02-02T02:15:00Z", "stage_status", "review_b", "completed", "DeepSeek 完成 v4 打分。"),
        (22, "2026-02-02T02:20:00Z", "stage_status", "review_c", "completed", "Qwen 完成 v4 打分。"),
        (23, "2026-02-02T02:25:00Z", "stage_status", "review_aggregation", "completed", "第 2 轮规则判定：0 个触发项，通过。"),
        (24, "2026-02-02T02:30:00Z", "human_gate", "expert_review", "waiting_human", "专家评分检查点：研究材料未包含专家评分。"),
        (25, "2026-02-02T02:30:00Z", "log", "finalize", "not_started", "等待专家评分，未进入最终入库。"),
    ]

    run = {
        "run_id": RUN_ID,
        "project_id": "sjt-university-echoes-2.0",
        "context_id": "university",
        "mode": "demo",
        "status": "waiting_human",
        "is_synthetic": False,
        "label": "校园回响：第七日重置 · 大学情境录制运行",
        "created_at": "2026-01-16T02:30:00Z",
        "config_snapshot": {
            "context_id": "university",
            "domain": "大五人格（BFI-2 构念）",
            "facet": "15 个 facet · 60 道情境题",
            "facet_definition": "基于 BFI-2 构念框架，覆盖五个维度下的 15 个 facet；每个 facet 4 道情境判断题。",
            "source_item_ids": [f"src-u{number:02d}" for number in range(1, 61)],
            "temperatures": {
                "item_transformation": None,
                "outline_generation": None,
                "story_generation": None,
            },
            "rubric_version": "ai-eval-0-1-2",
            "scoring_rule": SCORING_RULE,
            "max_revisions": 3,
            "batch_size": 60,
        },
        "provenance": {
            "source": "recorded",
            "generated_by": "研究材料整理脚本（frontend/scripts/build-recorded-run.py）",
            "disclaimer": "真实研究材料的录制回放；审查报告与 0/1/2 打分来自存档，不包含专家评分结论。",
            "exported_at": None,
        },
    }

    prompt_specs = [
        ("pv-item-transform", "item_transformation", "情景转化系统提示词", source / "01情景转换/prompt记录/情景转化系统提示词.txt"),
        ("pv-outline", "outline_generation", "大纲书写系统提示词", source / "03大纲设计/大纲书写系统提示词.txt"),
        ("pv-story", "story_generation", "故事生成系统提示词", source / "04故事生成/prompt记录/故事生成系统提示词.txt"),
        ("pv-review-agent", "review_agent", "剧情审查系统提示词", source / "05剧情评价/审查prompt/剧情评价系统提示词.txt"),
        ("pv-review", "review_a", "AI 评价系统提示词", source / "06AI评价打分/AI评价系统提示词.txt"),
        ("pv-review-b", "review_b", "AI 评价系统提示词", source / "06AI评价打分/AI评价系统提示词.txt"),
        ("pv-review-c", "review_c", "AI 评价系统提示词", source / "06AI评价打分/AI评价系统提示词.txt"),
        ("pv-revision", "auto_revision", "故事生成与修改系统提示词", source / "04故事生成/prompt记录/故事生成+修改系统提示词3.251.0.txt"),
    ]
    prompt_versions = []
    for prompt_id, stage_id, title, path in prompt_specs:
        if not path.exists():
            raise FileNotFoundError(path)
        prompt_versions.append(
            {
                "prompt_version_id": prompt_id,
                "stage_id": stage_id,
                "version": "研究存档（未编号）",
                "template": f"{title}\n\n{read_prompt(path)}",
                "created_at": "2026-01-16T00:00:00Z",
                "note": "提示词取自研究存档；运行时的确切版本未记录。",
            }
        )

    model_configs = [
        ("mc-transform", "item_transformation", "未记录", "未记录"),
        ("mc-outline", "outline_generation", "Google", "Gemini（具体版本未记录）"),
        ("mc-story", "story_generation", "Google", "Gemini（具体版本未记录）"),
        ("mc-review-agent", "review_agent", "Google / 未记录", "Gemini（具体版本未记录）"),
        ("mc-review-a", "review_a", "OpenAI", "GPT（具体版本未记录）"),
        ("mc-review-b", "review_b", "DeepSeek", "DeepSeek（具体版本未记录）"),
        ("mc-review-c", "review_c", "Alibaba", "Qwen（具体版本未记录）"),
        ("mc-revision", "auto_revision", "多模型与人工编辑", "具体模型未记录"),
        ("mc-expert", "expert_review", "human", "研究者（待评分）"),
    ]
    model_config_entries = [
        {
            "model_config_id": config_id,
            "stage_id": stage_id,
            "provider": provider,
            "model_id": model_id,
            "supported_params": [],
            "note": "研究材料未记录完整模型版本、温度与采样参数。",
        }
        for config_id, stage_id, provider, model_id in model_configs
    ]

    stages = []

    def add_stage(
        stage_id: str,
        attempt: int,
        revision: str | None,
        status: str,
        started_at: str | None,
        completed_at: str | None,
        output: dict | None = None,
        model: dict | None = None,
        prompt_version: str | None = None,
    ) -> None:
        stages.append(
            {
                "stage_execution_id": f"st-{stage_id}-{attempt}",
                "stage_id": stage_id,
                "attempt": attempt,
                "revision": revision,
                "status": status,
                "model": model,
                "prompt_version": prompt_version,
                "input": {},
                "output": output,
                "started_at": started_at,
                "completed_at": completed_at,
                "error": None,
            }
        )

    add_stage("intake", 1, None, "completed", "2026-01-16T02:00:00Z", "2026-01-16T02:05:00Z", {"source_item_count": 60})
    add_stage(
        "item_transformation",
        1,
        None,
        "completed",
        "2026-01-16T02:05:00Z",
        "2026-01-16T02:35:00Z",
        {
            "behavioral_targets": [
                {
                    "source_item_id": f"src-u{number:02d}",
                    "behavioral_description": f"{source_items[number - 1]['facet']} · {outline_descriptions[number - 1]}",
                }
                for number in range(1, 61)
            ]
        },
        prompt_version="pv-item-transform",
    )
    add_stage(
        "outline_generation",
        1,
        None,
        "completed",
        "2026-01-16T02:35:00Z",
        "2026-01-16T02:50:00Z",
        {"outline_version": OUTLINE_VERSION, "event_count": 60, "narrative_days": 7},
        prompt_version="pv-outline",
    )
    add_stage(
        "outline_confirmation",
        1,
        None,
        "completed",
        "2026-01-16T02:50:00Z",
        "2026-01-16T03:00:00Z",
        {"confirmed_version": OUTLINE_VERSION, "confirmed_by": "研究流程记录（本地研究者）"},
    )

    story_versions = [
        (1, "v1", "2026-01-16T03:05:00Z", "2026-01-16T03:20:00Z"),
        (2, "v2", "2026-01-16T04:13:00Z", "2026-01-16T04:20:00Z"),
        (3, "v3", "2026-01-23T06:48:00Z", "2026-01-23T07:00:00Z"),
        (4, "v4", "2026-02-02T01:22:00Z", "2026-02-02T01:30:00Z"),
    ]
    for number, version, started_at, completed_at in story_versions:
        add_stage(
            "story_generation",
            number,
            version,
            "completed",
            started_at,
            completed_at,
            {"story_version": version, "segments": 60, "items": 60},
            prompt_version="pv-story",
        )

    review_times = {
        1: ("2026-01-16T03:30:00Z", "2026-01-16T04:07:00Z"),
        2: ("2026-01-23T06:00:00Z", "2026-01-23T06:45:00Z"),
    }
    for report in review_reports:
        started_at, completed_at = review_times[report["attempt"]]
        add_stage(
            "review_agent",
            report["attempt"],
            report["reviewed_version"],
            "completed",
            started_at,
            completed_at,
            {
                "report_id": report["report_id"],
                "issue_count": len(report["issues"]),
                "checked_count": report["checked_count"],
            },
            model={
                "provider": "Google / 未记录",
                "model_id": report["model_label"],
                "parameters": {},
            },
            prompt_version="pv-review-agent",
        )

    scoring_labels = {
        "review_a": "GPT（具体版本未记录）",
        "review_b": "DeepSeek（具体版本未记录）",
        "review_c": "Qwen（具体版本未记录）",
    }
    scoring_times = {
        1: {
            "review_a": ("2026-01-29T05:30:00Z", "2026-01-29T05:45:00Z"),
            "review_b": ("2026-01-29T05:35:00Z", "2026-01-29T05:50:00Z"),
            "review_c": ("2026-01-29T05:55:00Z", "2026-01-29T06:10:00Z"),
        },
        2: {
            "review_a": ("2026-02-02T02:00:00Z", "2026-02-02T02:10:00Z"),
            "review_b": ("2026-02-02T02:02:00Z", "2026-02-02T02:15:00Z"),
            "review_c": ("2026-02-02T02:05:00Z", "2026-02-02T02:20:00Z"),
        },
    }
    for review in reviews:
        attempt = review["attempt"]
        slot = review["reviewer_slot"]
        started_at, completed_at = scoring_times[attempt][slot]
        prompt_version = {
            "review_a": "pv-review",
            "review_b": "pv-review-b",
            "review_c": "pv-review-c",
        }[slot]
        add_stage(
            slot,
            attempt,
            review["revision"],
            "completed",
            started_at,
            completed_at,
            {
                "review_id": review["review_id"],
                "trigger_count": len(round_triggers[attempt]),
                "recommendation": review["recommendation"],
            },
            model={
                "provider": scoring_labels[slot].split("（")[0],
                "model_id": scoring_labels[slot],
                "parameters": {},
            },
            prompt_version=prompt_version,
        )

    aggregation_times = {
        1: ("2026-01-29T06:15:00Z", "8 个触发项，退回定向修订。"),
        2: ("2026-02-02T02:25:00Z", "0 个触发项，通过并转专家评分。"),
    }
    for attempt, (at, decision) in aggregation_times.items():
        add_stage(
            "review_aggregation",
            attempt,
            None,
            "completed",
            at,
            at,
            {
                "scoring_rule": SCORING_RULE,
                "round": attempt,
                "trigger_count": len(round_triggers[attempt]),
                "trigger_item_ids": [
                    f"u-item-{number:02d}" for number in sorted(round_triggers[attempt])
                ],
                "decision": decision,
            },
        )

    revision_times = {
        1: "2026-01-16T04:13:00Z",
        2: "2026-01-23T06:48:00Z",
        3: "2026-02-02T01:22:00Z",
    }
    for number, revision in enumerate(revisions, start=1):
        add_stage(
            "auto_revision",
            number,
            f"{revision['from_version']}→{revision['to_version']}",
            "completed",
            None,
            revision_times[number],
            {
                "revision_id": revision["revision_id"],
                "changed_count": len(revision["changed_item_ids"]),
            },
            prompt_version="pv-revision",
        )

    add_stage(
        "expert_review",
        1,
        None,
        "waiting_human",
        None,
        None,
        {"decision": "pending", "note": "研究材料未包含专家评分记录。"},
    )

    story = {
        "story_id": "story-univ-echoes-v4",
        "version": STORY_VERSION,
        "outline_version": OUTLINE_VERSION,
        "context_id": "university",
        "title": "校园回响：第七日重置",
        "text": "\n\n".join(final_story[number]["text"] for number in range(1, 61)),
        "segments": [
            {"event_id": f"oe-{number:02d}", "text": final_story[number]["text"]}
            for number in range(1, 61)
        ],
    }

    return {
        "schema_version": "sjt-run/1",
        "run": run,
        "source_items": source_records,
        "prompt_versions": prompt_versions,
        "model_configs": model_config_entries,
        "stages": stages,
        "events": [
            {
                "seq": seq,
                "event_id": f"evt-{seq:03d}",
                "at": at,
                "type": event_type,
                "stage_id": stage_id,
                "status": status,
                "message": message,
            }
            for seq, at, event_type, stage_id, status, message in events
        ],
        "outline": {
            "outline_version": OUTLINE_VERSION,
            "confirmed_version": OUTLINE_VERSION,
            "confirmed_by": "研究流程记录（本地研究者）",
            "confirmed_at": "2026-01-16T03:00:00Z",
            "events": outline_events,
        },
        "story": story,
        "items": items,
        "review_reports": review_reports,
        "reviews": reviews,
        "scoring_examples": scoring_examples,
        "revisions": revisions,
        "expert_decision": {
            "decision": "pending",
            "item_version_id": None,
            "reviewer": None,
            "comment": None,
            "decided_at": None,
        },
        "limitations": [
            "打分智能体的条目分取存档剧情层的 0/1/2 评分；问句与选项层由审查智能体逐项核查。",
            "触发规则：任一维度出现 0 分，或同一维度有 2 个及以上智能体给 1 分，才退回修改。",
            "最终版取自 07 专家评审最终用版，选项分值 A=1、B=2、C=3、D=4。",
            "专家评分尚未记录，expert_decision 保持 pending。",
            "模型名称来自研究材料；具体版本、温度与采样参数未记录。",
            "事件时间依据研究材料顺序重建，不是原始系统日志时间。",
            "AI 打分不等同于心理测量效度或专家内容效度判断。",
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument(
        "--out",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "data/runs/university-echoes-2.0.json",
    )
    args = parser.parse_args()
    payload = build(args)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(
        f"wrote {args.out} ({args.out.stat().st_size:,} bytes): "
        f"{len(payload['items'])} items, {len(payload['review_reports'])} review reports, "
        f"{len(payload['reviews'])} scoring records, {len(payload['revisions'])} revisions, "
        f"{len(payload['events'])} events"
    )


if __name__ == "__main__":
    main()
